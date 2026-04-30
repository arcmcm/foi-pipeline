#!/usr/bin/env python3
"""
FOI research prototype (homelessness-focused)

What it does
- Reads seeds from seeds.txt (local file paths OR URLs)
- Normalizes PDF / DOCX / HTML / TXT
- For homelessness FOI responses like the Harvey Glasgow FOI:
  - Extracts text from DOCX tables
  - Splits one big table into “question sections”
  - Infers a metric per section
  - Extracts plausible count numbers
  - Filters out year/date noise and Section 12 / £600 cost-limit noise
  - Deduplicates values within each section
- Writes data/report.json

IMPORTANT
- This file must start with: #!/usr/bin/env python3
- Do NOT paste terminal commands like: cd ... or cat > ... into this file
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pdfplumber
import requests
from bs4 import BeautifulSoup

try:
    import docx  # python-docx
except Exception:
    docx = None


# -----------------------------
# Paths / config
# -----------------------------
DATA_DIR = Path("data")
REPORT_PATH = DATA_DIR / "report.json"
SEEDS_PATH = Path("seeds.txt")

DEFAULT_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux aarch64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/123 Safari/537.36"
    )
}

MODE = "homelessness"

STAT_KEYWORDS = [
    "homeless",
    "temporary accommodation",
    "applications",
    "single person applications",
    "family applications",
    "registered as homeless",
    "temporary accommodation units",
    "four hurdles",
]

# Boundary-safe number matcher:
# - matches "10,501" and "3967"
# - avoids matching "200" as part of "2006"
NUMBER_RE = re.compile(
    r"(?<!\d)(?P<value>(?:<\s*)?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)(?!\d)"
)

DATE_RANGE_RE = re.compile(
    r"(\d{1,2}(?:st|nd|rd|th)?\s+[A-Za-z]+\s+\d{4}\s*[–-]\s*"
    r"\d{1,2}(?:st|nd|rd|th)?\s+[A-Za-z]+\s+\d{4})",
    flags=re.I,
)
AS_AT_DATE_RE = re.compile(
    r"(?:as at|as of)\s+(\d{1,2}(?:st|nd|rd|th)?\s+[A-Za-z]+\s+\d{4})",
    flags=re.I,
)


# -----------------------------
# Data models
# -----------------------------
@dataclass
class TextBlock:
    locator: str
    raw_content: str


@dataclass
class DocumentRecord:
    url: str
    file_type: str
    sha256: str
    retrieved_at: str
    text_blocks: List[TextBlock]


@dataclass
class ClaimRecord:
    metric: str
    value_raw: str
    context_text: str
    document_url: str
    locator: str
    date_hint: Optional[str] = None
    question_id: Optional[str] = None


# -----------------------------
# Utilities
# -----------------------------
def now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def sha256_bytes(b: bytes) -> str:
    import hashlib as _hashlib

    return _hashlib.sha256(b).hexdigest()


def read_seeds(path: Path) -> List[str]:
    if not path.exists():
        return []
    out: List[str] = []
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        out.append(s)
    return out


def is_url(s: str) -> bool:
    return s.startswith("http://") or s.startswith("https://")


def guess_file_type_from_url_or_path(s: str) -> str:
    low = s.lower()
    if low.endswith(".pdf"):
        return "pdf"
    if low.endswith(".docx"):
        return "docx"
    if low.endswith(".html") or low.endswith(".htm"):
        return "html"
    if low.endswith(".txt"):
        return "txt"
    return "unknown"


def fetch_bytes(seed: str) -> Tuple[bytes, str]:
    if is_url(seed):
        r = requests.get(seed, headers=DEFAULT_HEADERS, timeout=60)
        r.raise_for_status()
        return r.content, r.url

    p = Path(seed)
    if not p.exists():
        raise FileNotFoundError(seed)
    return p.read_bytes(), str(p.resolve())


# -----------------------------
# Normalizers
# -----------------------------
def normalize_pdf(content: bytes, url: str) -> List[TextBlock]:
    blocks: List[TextBlock] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            txt = page.extract_text() or ""
            if txt.strip():
                blocks.append(TextBlock(locator=f"page {i}", raw_content=txt))
    return blocks


def normalize_html(content: bytes, url: str) -> List[TextBlock]:
    html = content.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    text = soup.get_text("\n")
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    cleaned = "\n".join(lines)
    if not cleaned.strip():
        return []
    return [TextBlock(locator="html", raw_content=cleaned)]


def normalize_txt(content: bytes, url: str) -> List[TextBlock]:
    txt = content.decode("utf-8", errors="ignore").strip()
    if not txt:
        return []
    return [TextBlock(locator="text", raw_content=txt)]


def normalize_docx(content: bytes, url: str) -> List[TextBlock]:
    if docx is None:
        raise RuntimeError("python-docx not installed. Install with: pip install python-docx")

    document = docx.Document(io.BytesIO(content))
    blocks: List[TextBlock] = []

    # Paragraphs (kept, but extractor will ignore in homelessness mode)
    paras = [(p.text or "").strip() for p in document.paragraphs]
    paras = [p for p in paras if p]
    if paras:
        blocks.append(TextBlock(locator="docx paragraphs", raw_content="\n".join(paras)))

    # Tables (main data source)
    for t_index, table in enumerate(document.tables, start=1):
        rows_text: List[str] = []
        for row in table.rows:
            cells = [" ".join((cell.text or "").split()) for cell in row.cells]
            rows_text.append("\t".join(cells))
        table_txt = "\n".join([row for row in rows_text if row.strip()])
        if table_txt.strip():
            blocks.append(TextBlock(locator=f"docx table {t_index}", raw_content=table_txt))

    return blocks


def normalize_document(seed: str, content: bytes, final_url: str) -> DocumentRecord:
    ft = guess_file_type_from_url_or_path(seed)
    if ft == "unknown":
        ft = guess_file_type_from_url_or_path(final_url)

    if ft == "pdf":
        blocks = normalize_pdf(content, final_url)
    elif ft == "docx":
        blocks = normalize_docx(content, final_url)
    elif ft == "html":
        blocks = normalize_html(content, final_url)
    elif ft == "txt":
        blocks = normalize_txt(content, final_url)
    else:
        # best effort
        blocks = normalize_html(content, final_url)
        if not blocks:
            blocks = normalize_txt(content, final_url)

    return DocumentRecord(
        url=final_url,
        file_type=ft,
        sha256=sha256_bytes(content),
        retrieved_at=now_iso(),
        text_blocks=blocks,
    )


# -----------------------------
# Extraction helpers
# -----------------------------
def extract_date_hint(text: str) -> Optional[str]:
    m = AS_AT_DATE_RE.search(text)
    if m:
        return m.group(1)
    m = DATE_RANGE_RE.search(text)
    if m:
        return m.group(1)
    return None


def parse_int_like(value_raw: str) -> Optional[int]:
    v = value_raw.replace(",", "").strip()
    if v.startswith("<"):
        v = v.lstrip("<").strip()
    try:
        return int(float(v))
    except ValueError:
        return None


def is_plausible_homeless_count(n: int) -> bool:
    # Drop years and small list numbers (Q1..Q6, 31st March, etc.)
    if 1900 <= n <= 2100:
        return False
    if n < 50:
        return False
    return True


def infer_metric_from_text(text: str) -> Optional[str]:
    low = text.lower()

    # Order matters: more specific before less specific
    if "single person applications" in low and "family applications" in low:
        return "homeless_applications_by_household_type"
    if "four hurdles" in low:
        return "homeless_applications_passed_four_hurdles"
    if "registered as homeless" in low:
        return "households_registered_homeless"
    if "temporary accommodation units" in low:
        return "temporary_accommodation_units"
    if 'how many "homeless applications"' in low or "how many homeless applications" in low:
        return "homeless_applications"

    return None


def split_table_into_sections(table_text: str) -> List[str]:
    """
    Splits a flattened DOCX table into sections, using “header-ish” lines.
    This works better than relying on clean Q1/Q2 formatting because the
    Word table often repeats header text.
    """
    lines = [ln.strip() for ln in table_text.splitlines() if ln.strip()]
    if not lines:
        return []

    sections: List[List[str]] = []
    current: List[str] = []

    def is_header(line: str) -> bool:
        low = line.lower()
        if low.startswith("number of "):
            return True
        # In this FOI the question text contains '?' but may not END with '?'
        if "how many" in low and "?" in line:
            return True
        return False

    for line in lines:
        if is_header(line) and current:
            sections.append(current)
            current = [line]
        else:
            current.append(line)

    if current:
        sections.append(current)

    return ["\n".join(sec) for sec in sections if sec]


def extract_claims_homelessness(doc: DocumentRecord) -> List[ClaimRecord]:
    claims: List[ClaimRecord] = []

    for block in doc.text_blocks:
        # For this FOI, the real counts are in DOCX tables.
        if not block.locator.startswith("docx table"):
            continue

        text = block.raw_content
        low = text.lower()

        if not any(k in low for k in STAT_KEYWORDS):
            continue

        sections = split_table_into_sections(text) or [text]

        for section_text in sections:
            metric = infer_metric_from_text(section_text)
            if not metric:
                continue

            date_hint = extract_date_hint(section_text)

            seen: set[Tuple[str, str]] = set()

            for match in NUMBER_RE.finditer(section_text):
                value_raw = match.group("value").strip()
                n = parse_int_like(value_raw)
                if n is None:
                    continue
                if not is_plausible_homeless_count(n):
                    continue

                start = max(0, match.start() - 120)
                end = min(len(section_text), match.end() + 200)
                context = section_text[start:end].replace("\n", " ").strip()
                context_low = context.lower()

                # Filter the cost-limit noise section (Section 12 / £600)
                if "section 12" in context_low or "upper limit" in context_low or "£600" in context:
                    continue

                key = (metric, value_raw)
                if key in seen:
                    continue
                seen.add(key)

                claims.append(
                    ClaimRecord(
                        metric=metric,
                        value_raw=value_raw,
                        context_text=context,
                        document_url=doc.url,
                        locator=block.locator,
                        date_hint=date_hint,
                        question_id=None,
                    )
                )

    return claims


def extract_claims_from_document(doc: DocumentRecord) -> List[ClaimRecord]:
    if MODE == "homelessness":
        return extract_claims_homelessness(doc)
    return []


# -----------------------------
# Pipeline
# -----------------------------
def run_pipeline() -> Dict[str, int]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    seeds = read_seeds(SEEDS_PATH)
    documents: List[DocumentRecord] = []
    claims: List[ClaimRecord] = []

    for seed in seeds:
        try:
            content, final_url = fetch_bytes(seed)
            doc = normalize_document(seed, content, final_url)
        except Exception as e:
            print(f"[WARN] seed failed/normalize failed: {seed} ({e})", file=sys.stderr)
            continue

        documents.append(doc)
        claims.extend(extract_claims_from_document(doc))

    report: Dict[str, object] = {
        "generated_at": now_iso(),
        "mode": MODE,
        "documents": [
            {
                "url": d.url,
                "file_type": d.file_type,
                "sha256": d.sha256,
                "retrieved_at": d.retrieved_at,
                "text_blocks": [asdict(b) for b in d.text_blocks],
            }
            for d in documents
        ],
        "claims": [asdict(c) for c in claims],
        "findings": [],
    }

    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    return {
        "document_count": len(documents),
        "claim_count": len(claims),
        "finding_count": 0,
    }


def main() -> None:
    summary = run_pipeline()
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
